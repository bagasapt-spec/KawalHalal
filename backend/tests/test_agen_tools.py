"""Tes deterministik untuk tool agen, gerbang konfirmasi, dan sinkron ke mock SIHALAL (tanpa model)."""
import os
import sys
from types import SimpleNamespace

import pytest

BACKEND_DIR = os.path.abspath(os.path.join(os.path.dirname(__file__), ".."))
if BACKEND_DIR not in sys.path:
    sys.path.insert(0, BACKEND_DIR)

import sihalal_mock as sm  # noqa: E402
from kawal_agent import gates, sinkron, templates, tools as T  # noqa: E402


class Ctx:
    """Pengganti ToolContext/CallbackContext: hanya atribut yang dipakai tool."""

    def __init__(self, user_id="+6281200000001"):
        self.state = {}
        self.user_id = user_id
        self.invocation_id = "inv-0"
        self.user_content = None
        self._n = 0

    def giliran(self, teks: str):
        """Mulai giliran baru: jalankan callback gate dengan pesan pengguna sungguhan."""
        self._n += 1
        self.invocation_id = f"inv-{self._n}"
        self.user_content = SimpleNamespace(parts=[SimpleNamespace(text=teks)])
        gates.proses_giliran_baru(self)
        return self


@pytest.fixture(autouse=True)
def bersih(tmp_path, monkeypatch):
    monkeypatch.setattr(T, "GENERATED_DIR", tmp_path)
    sm.muat_seed()
    sm._akun.clear()
    yield


def siapkan_draf(c: Ctx, kategori="makanan") -> Ctx:
    T.pilih_kategori(kategori, c)
    r = T.daftarkan_akun_sihalal("Bakso Pak Joko", "Joko Susanto", "joko@example.com", "081234567890", "9120001234567",
                                 "Jl. Mawar 1", "Kota Bandung", "Jawa Barat", "40135", "mikro", c)
    assert r["status"] == "terdaftar", r
    T.tentukan_jalur_sertifikasi(1, True, 500_000_000, c)
    r = T.buat_draf_pengajuan("self-declare", c)
    assert r["status"] == "dibuat", r
    return c


def pengajuan(c: Ctx) -> dict:
    return sm._cari(c.state["pengajuan_id"])


# ------------------------------------------------------------ gerbang
@pytest.mark.parametrize("teks,hasil", [
    ("YA", "ya"), ("iya", "ya"), ("Ya, masukkan saja", "ya"), ("oke lanjut", "ya"), ("sip", "ya"),
    ("tidak", "tidak"), ("nggak usah", "tidak"), ("batal", "tidak"),
    ("ya tapi ganti nama bahannya dulu", None), ("hmm", None), ("", None), ("belum yakin", None),
    ("ya ya ya ya ya ya ya ya ya ya", None),
])
def test_baca_jawaban(teks, hasil):
    assert gates.baca_jawaban(teks) == hasil


def test_model_tidak_bisa_menyetujui_sendiri():
    c = siapkan_draf(Ctx().giliran("halo"))
    r = T.tambah_bahan("Gelatin", "bahan", "", "", "", "", "", c)
    assert r["verdict"] == "perlu_tindak_lanjut" and r["masuk_sihalal"] is False
    T.minta_konfirmasi("masukkan_bahan_bermasalah", "Gelatin belum ada sertifikat. Tetap masukkan?", c)
    # di giliran yang sama model memaksa sinkron: tetap ditahan
    T.sinkron_ke_sihalal(c)
    assert not pengajuan(c)["bahan"]["daftarBahan"]
    # jawaban ambigu tidak menyetujui
    c.giliran("ya tapi nanti dulu")
    T.sinkron_ke_sihalal(c)
    assert not pengajuan(c)["bahan"]["daftarBahan"]
    # jawaban jelas menyetujui
    c.giliran("ya")
    assert c.state["keputusan_terakhir"]["keputusan"] == "disetujui"
    r = T.sinkron_ke_sihalal(c)
    b = pengajuan(c)["bahan"]["daftarBahan"]
    assert [x["nama"] for x in b] == ["Gelatin"] and b[0]["bukti"] == "kurang" and b[0]["konfirmasi"] == "terkonfirmasi"
    assert r["bahan_ditahan_menunggu_konfirmasi"] == []


def test_penolakan_menahan_bahan():
    c = siapkan_draf(Ctx().giliran("halo"))
    T.tambah_bahan("Gelatin", "bahan", "", "", "", "", "", c)
    T.minta_konfirmasi("masukkan_bahan_bermasalah", "Tetap masukkan?", c)
    c.giliran("tidak, saya cari sertifikatnya dulu")
    assert c.state["keputusan_terakhir"]["keputusan"] == "ditolak"
    T.sinkron_ke_sihalal(c)
    assert not pengajuan(c)["bahan"]["daftarBahan"]


# ------------------------------------------------------------ verifikasi dan aturan masuk
def test_bahan_bersih_langsung_masuk_tanpa_tanya():
    c = siapkan_draf(Ctx().giliran("halo"))
    r = T.tambah_bahan("Kecap manis", "bahan", "Kecap Cap Layar", "ID1230000100101", "Toko Sinar", "", "", c)
    assert r["verdict"] == "bersih" and r["masuk_sihalal"] is True
    b = pengajuan(c)["bahan"]["daftarBahan"][0]
    assert b["bukti"] == "ada" and b["noSertifikat"] == "ID1230000100101" and b["sumber"]["oleh"] == "kawalhalal"
    r = T.tambah_bahan("Tepung tapioka", "bahan", "", "", "Pasar", "", "", c)
    assert r["verdict"] == "bersih" and r["masuk_sihalal"] is True


@pytest.mark.parametrize("nomor,verdict,bukti", [
    ("ID1230000100106", "perlu_tindak_lanjut", "kurang"),       # kedaluwarsa
    ("ID9999999999999", "perlu_tindak_lanjut", "tidak-jelas"),  # tidak terdaftar
    ("ID1230000100105", "perlu_tindak_lanjut", "tidak-jelas"),  # nomor milik produk lain
])
def test_sertifikat_bermasalah(nomor, verdict, bukti):
    c = siapkan_draf(Ctx().giliran("halo"))
    r = T.tambah_bahan("Santan kelapa", "bahan", "Santan Kelapa Mas", nomor, "", "", "", c)
    assert (r["verdict"], r["status_bukti"]) == (verdict, bukti)
    assert r["masuk_sihalal"] is False and not pengajuan(c)["bahan"]["daftarBahan"]


def test_daging_ditelusuri_dari_lapak():
    c = siapkan_draf(Ctx().giliran("halo"))
    r = T.tambah_bahan("Daging sapi", "bahan", "", "", "Lapak Daging Barokah Pak Slamet", "sapi", "", c)
    assert r["verdict"] == "bersih" and any("RPH Cakung" in x for x in r["catatan_penelusuran"])
    r = T.tambah_bahan("Daging kambing", "bahan", "", "", "Lapak Daging Kambing Haji Ujang", "kambing", "", c)
    assert r["verdict"] == "perlu_tindak_lanjut"
    r = T.tambah_bahan("Ayam", "bahan", "", "", "Lapak Tidak Dikenal", "ayam", "", c)
    assert r["verdict"] == "perlu_tindak_lanjut"
    r = T.tambah_bahan("Ayam", "bahan", "", "ID3175000200003", "Lapak Ayam Mitra Pak Dedi", "ayam", "", c)
    assert r["status_bukti"] == "kurang"  # RPU kedaluwarsa


def test_babi_ditolak_dan_tidak_masuk():
    c = siapkan_draf(Ctx().giliran("halo"))
    r = T.tambah_bahan("Lemak babi", "bahan", "", "", "", "", "", c)
    assert r["status"] == "ditolak_tidak_dapat_disertifikasi"
    assert not _dok(c)["bahan"] and not pengajuan(c)["bahan"]["daftarBahan"]


def _dok(c):
    return c.state.get("dok") or {"bahan": []}


def test_foto_bukti_harus_ada():
    c = siapkan_draf(Ctx().giliran("halo"))
    r = T.tambah_bahan("Kecap manis", "bahan", "Kecap Cap Layar", "ID1230000100101", "", "", "karangan.jpg", c)
    assert any("tidak ada" in x for x in r["catatan_penelusuran"])
    assert "fotoSumber" not in pengajuan(c)["bahan"]["daftarBahan"][0]


def test_bahan_tercatat_sebelum_draf_masuk_setelah_draf():
    c = Ctx().giliran("halo")
    T.pilih_kategori("minuman", c)
    r = T.tambah_bahan("Teh celup", "bahan", "Teh Melati Asri", "ID1230000100107", "", "", "", c)
    assert r["verdict"] == "bersih" and r["masuk_sihalal"] is False
    T.daftarkan_akun_sihalal("Es Teh Bu Ani", "Ani", "ani@example.com", "08123456789", "", "Jl. A", "Bandung", "Jawa Barat", "40111", "mikro", c)
    T.buat_draf_pengajuan("self-declare", c)
    assert [b["nama"] for b in pengajuan(c)["bahan"]["daftarBahan"]] == ["Teh celup"]


# ------------------------------------------------------------ akun dan validasi
def test_pendaftaran_akun_memeriksa_data():
    c = Ctx().giliran("halo")
    r = T.daftarkan_akun_sihalal("X", "Y", "bukan-email", "12", "123", "Jl", "Kota", "Prov", "1", "raksasa", c)
    assert r["status"] == "data_tidak_valid" and len(r["masalah"]) >= 5
    assert not c.state.get("akun")
    r = T.daftarkan_akun_sihalal("Warung Siti", "Siti", "siti@example.com", "+62 812-3456-7890", "", "Jl. B", "Bandung", "Jawa Barat", "40111", "mikro", c)
    assert r["status"] == "terdaftar" and "NIB" in r["pesan"]
    assert sm.cek_masuk(r["username"], r["password_sementara"]) and not sm.cek_masuk(r["username"], "salah")


def test_pesan_terstruktur_tanggal_dan_jumlah_salah():
    c = siapkan_draf(Ctx().giliran("halo"))
    T.tambah_bahan("Tepung tapioka", "bahan", "", "", "", "", "", c)
    r = T.catat_pembelian_bahan("Tepung tapioka", 0, "kg", "31-10-2026", "", "2026-02-30", c)
    assert r["status"] == "data_tidak_valid" and len(r["masalah"]) == 4
    r = T.catat_pembelian_bahan("Bahan tak dikenal", 1, "kg", "2026-10-01", "Dapur", "2026-10-01", c)
    assert r["status"] == "data_tidak_valid"


# ------------------------------------------------------------ end to end sampai kirim
def lengkapi_semua(c: Ctx):
    T.isi_data_pengajuan("nomorSurat", "001/BPJ/X/2026", c)
    T.catat_produk("Bakso Sapi", "Bakso Pak Joko", "Produk olahan daging", c)
    assert T.tambah_bahan("Daging sapi", "bahan", "", "", "Lapak Daging Barokah Pak Slamet", "sapi", "", c)["masuk_sihalal"]
    assert T.tambah_bahan("Kecap manis", "bahan", "Kecap Cap Layar", "ID1230000100101", "Toko A", "", "", c)["masuk_sihalal"]
    assert T.tambah_bahan("Tepung tapioka", "bahan", "", "", "Pasar", "", "", c)["masuk_sihalal"]
    for n in ("Daging sapi", "Kecap manis", "Tepung tapioka"):
        assert T.catat_pembelian_bahan(n, 2, "kg", "2026-10-20", "Dapur", "2026-10-20", c)["status"] == "ok"
    assert T.catat_alur_proses("Pembuatan bakso", "Giling daging; Campur tapioka; Bentuk bulat; Rebus", c)["status"] == "ok"
    assert T.catat_catatan_produksi("penyimpanan_bahan", "Daging sapi", 5, "2026-10-20", "", c)["status"] == "ok"
    assert T.catat_catatan_produksi("hasil_produksi", "Bakso Sapi", 100, "2026-10-21", "", c)["status"] == "ok"
    assert T.catat_catatan_produksi("distribusi", "Bakso Sapi", 50, "2026-10-21", "Warung sendiri", c)["status"] == "ok"
    T.tambah_anggota_komitmen("Joko Susanto", "Pemilik", "ketua", c)
    c.state["berkas_masuk"] = [{"nama": "denah.jpg", "mime": "image/jpeg", "url": "u1"},
                               {"nama": "ttd-pj.jpg", "mime": "image/jpeg", "url": "u2"},
                               {"nama": "ttd-penyelia.jpg", "mime": "image/jpeg", "url": "u3"}]
    assert T.lampirkan_berkas("layout", "denah.jpg", c)["status"] == "ok"
    assert T.lampirkan_berkas("tanda_tangan_pj", "ttd-pj.jpg", c)["status"] == "ok"
    assert T.lampirkan_berkas("tanda_tangan_penyelia", "ttd-penyelia.jpg", c)["status"] == "ok"


def setujui_pernyataan(c: Ctx):
    for j in T.TEKS_PERNYATAAN:
        r = T.setujui_pernyataan(j, c)
        assert r["status"] == "menunggu_pengguna" and r["teks_pernyataan"]
        c.giliran("YA")
        assert T.setujui_pernyataan(j, c)["status"] == "tercatat"


def test_dari_nol_sampai_terkirim():
    c = siapkan_draf(Ctx().giliran("mulai"))
    lengkapi_semua(c)

    cek = T.periksa_pengajuan(c)
    assert not cek["siap_kirim"] and any("persetujuan kebijakan halal" in x for x in cek["harus_diperbaiki"])
    # belum siap: kirim ditolak, tidak meminta konfirmasi dulu
    assert T.kirim_pengajuan_sihalal(c)["status"] == "belum_siap"

    setujui_pernyataan(c)
    cek = T.periksa_pengajuan(c)
    assert cek["siap_kirim"], cek
    assert pengajuan(c)["status"] == "draf"

    r = T.kirim_pengajuan_sihalal(c)  # minta konfirmasi
    assert r["status"] == "menunggu_pengguna" and pengajuan(c)["status"] == "draf"
    assert T.kirim_pengajuan_sihalal(c)["status"] == "menunggu_pengguna"  # model tidak bisa memaksa
    c.giliran("ya kirim")
    r = T.kirim_pengajuan_sihalal(c)
    assert r["status"] == "terkirim" and r["nomor_daftar"].startswith("SH2026-1-")
    p = pengajuan(c)
    assert p["status"] == "terkirim" and any(a["aksi"] == "Mengirim pengajuan" and a["pelaku"] == "kawalhalal" for a in p["audit"])
    # sesudah terkirim tidak bisa diubah
    assert T.isi_data_pengajuan("merek", "Lain", c)["status"] == "galat"


def test_persetujuan_kirim_basi_jika_data_berubah():
    c = siapkan_draf(Ctx().giliran("mulai"))
    lengkapi_semua(c)
    setujui_pernyataan(c)
    T.kirim_pengajuan_sihalal(c)
    c.giliran("ya")
    # data berubah setelah menyetujui, sebelum kirim
    T.catat_alur_proses("Pengemasan", "Dinginkan; Kemas", c)
    r = T.kirim_pengajuan_sihalal(c)
    assert r["status"] == "menunggu_pengguna", "persetujuan lama tidak boleh dipakai untuk data yang sudah berubah"


def test_kirim_diblokir_bila_ada_bahan_menunggu():
    c = siapkan_draf(Ctx().giliran("mulai"))
    lengkapi_semua(c)
    setujui_pernyataan(c)
    T.tambah_bahan("Gelatin", "bahan", "", "", "", "", "", c)
    r = T.kirim_pengajuan_sihalal(c)
    assert r["status"] == "belum_siap" and any("Gelatin" in x for x in r["harus_diperbaiki"])


def test_bahan_bercelah_jadi_peringatan_saat_kirim():
    c = siapkan_draf(Ctx().giliran("mulai"))
    lengkapi_semua(c)
    setujui_pernyataan(c)
    T.tambah_bahan("Gelatin", "bahan", "", "", "", "", "", c)
    T.minta_konfirmasi("masukkan_bahan_bermasalah", "Tetap?", c)
    c.giliran("ya")
    T.sinkron_ke_sihalal(c)
    T.catat_pembelian_bahan("Gelatin", 1, "kg", "2026-10-20", "Dapur", "2026-10-20", c)
    cek = T.periksa_pengajuan(c)
    assert cek["siap_kirim"] and any("Gelatin" in x for x in cek["peringatan"])
    assert "Gelatin" in T.kirim_pengajuan_sihalal(c)["ringkasan"]


# ------------------------------------------------------------ xlsx
def test_xlsx_belum_diisi_dan_impor_balik(tmp_path):
    c = siapkan_draf(Ctx().giliran("mulai"), "makanan")
    T.catat_produk("Bakso Sapi", "Bakso Pak Joko", "", c)  # rincian kosong -> wajib di xlsx
    T.tambah_bahan("Tepung tapioka", "bahan", "", "", "", "", "", c)
    r = T.buat_berkas_xlsx(c)
    assert r["status"] == "dibuat" and r["sel_belum_diisi"] >= 1
    path = T.folder_pengguna(c.user_id) / r["nama_berkas"]
    from openpyxl import load_workbook
    nilai = [cell.value for row in load_workbook(path)["Produk"].iter_rows(min_row=2) for cell in row]
    assert templates.BELUM_DIISI in nilai
    assert "Tepung tapioka" in [cell.value for row in load_workbook(path)["Bahan"].iter_rows(min_row=2) for cell in row]
    assert "Produk: baris 1" not in r  # hanya memastikan hasil dict, bukan teks
    # pengguna mengisi sendiri lewat xlsx lalu mengirim balik
    from openpyxl import Workbook
    wb = load_workbook(path)
    wb["Bahan"].cell(row=2, column=5, value="Pasar Ciroyom")
    wb["Bahan"].append(["Kecap manis", "bahan", "Kecap Cap Layar", "ID1230000100101", "Toko A", "dibuat pengguna sendiri"])
    isi = T.folder_pengguna(c.user_id) / "isi-sendiri.xlsx"
    wb.save(isi)
    c.state["berkas_masuk"] = [{"nama": "isi-sendiri.xlsx", "mime": "x", "path": str(isi), "url": "u"}]
    h = T.impor_berkas_xlsx("isi-sendiri.xlsx", c)
    assert h["status"] == "diimpor"
    nama = {b["nama"]: b for b in pengajuan(c)["bahan"]["daftarBahan"]}
    assert set(nama) == {"Tepung tapioka", "Kecap manis"}
    # kolom status buatan pengguna diabaikan, status dihitung ulang dari registri
    assert nama["Kecap manis"]["bukti"] == "ada"


def test_impor_berkas_tidak_ada_atau_rusak(tmp_path):
    c = siapkan_draf(Ctx().giliran("mulai"))
    assert T.impor_berkas_xlsx("hantu.xlsx", c)["status"] == "berkas_tidak_ada"
    rusak = tmp_path / "rusak.xlsx"
    rusak.write_bytes(b"bukan xlsx")
    c.state["berkas_masuk"] = [{"nama": "rusak.xlsx", "path": str(rusak)}]
    assert T.impor_berkas_xlsx("rusak.xlsx", c)["status"] == "tidak_terbaca"


def test_skema_kategori_berbeda():
    assert "asal" in [k["kunci"] for k in templates.skema_kolom("kosmetik")["bahan"]]
    assert "asal" not in [k["kunci"] for k in templates.skema_kolom("makanan")["bahan"]]
    assert templates.skema_kolom("jasa")["produk"][0]["label"] == "Nama Layanan"


# ------------------------------------------------------------ sinkron
def test_sinkron_idempoten_dan_tidak_menimpa_isian_web():
    c = siapkan_draf(Ctx().giliran("mulai"))
    T.tambah_bahan("Tepung tapioka", "bahan", "", "", "", "", "", c)
    # pengguna menambah bahan sendiri lewat web
    with sm._kunci:
        p = sm._cari(c.state["pengajuan_id"])
        isi = p["bahan"]
        isi["daftarBahan"].append({"id": "b-web", "tipe": "tidak-bersertifikat", "jenis": "bahan", "nama": "Garam", "bukti": "ada", "konfirmasi": "terkonfirmasi"})
        sm.ubah_bagian(p["id"], "bahan", isi, pelaku="pengguna")
    n_audit = len(pengajuan(c)["audit"])
    r1 = T.sinkron_ke_sihalal(c)
    r2 = T.sinkron_ke_sihalal(c)
    assert r1["bagian_diubah"] == [] and r2["bagian_diubah"] == []
    assert len(pengajuan(c)["audit"]) == n_audit, "sinkron tanpa perubahan tidak boleh menulis audit"
    assert {b["nama"] for b in pengajuan(c)["bahan"]["daftarBahan"]} == {"Tepung tapioka", "Garam"}


def test_rincian_kurang_sejalan_dengan_hitung_kelengkapan():
    c = siapkan_draf(Ctx().giliran("mulai"))
    for tahap in range(3):
        p = pengajuan(c)
        total = sum(v["total"] for v in sm.hitung_kelengkapan(p).values())
        selesai = sum(v["done"] for v in sm.hitung_kelengkapan(p).values())
        assert len(sinkron.rincian_kurang(p)) == total - selesai
        if tahap == 0:
            lengkapi_semua(c)
        elif tahap == 1:
            setujui_pernyataan(c)


def test_lihat_progres_dan_pilihan_kategori():
    c = Ctx().giliran("halo")
    r = T.lihat_progres(c)
    assert len(r["pilihan_kategori"]) == 5 and r["langkah_berikutnya"] == "Pilih kategori produk"
    assert T.pilih_kategori("2", c)["kategori"] == "minuman"
    assert T.pilih_kategori("kosmetik", c)["kategori"] == "kosmetik"
    assert T.pilih_kategori("pesawat", c)["status"] == "tidak_dikenal"



def test_lampirkan_berkas_fallback_cli_tanpa_upload():
    """Di CLI/terminal, penyebutan nama berkas seperti denah_dapur.jpg atau dokumen_komitmen_tanda_tangan.jpg langsung diterima tanpa upload biner."""
    c = siapkan_draf(Ctx().giliran("halo"))
    r1 = T.lampirkan_berkas("denah", "denah_dapur.jpg", c)
    assert r1["status"] == "ok" and r1["berkas"] == "denah_dapur.jpg"
    r2 = T.lampirkan_berkas("komitmen", "dokumen_komitmen_tanda_tangan.jpg", c)
    assert r2["status"] == "ok" and r2["berkas"] == "dokumen_komitmen_tanda_tangan.jpg"
    p = pengajuan(c)
    assert p["proses"]["layout"][0]["berkas"]["nama"] == "denah_dapur.jpg"
    assert p["pemantauan"]["tandaTangan"][0]["penanggungJawab"]["nama"] == "dokumen_komitmen_tanda_tangan.jpg"
    assert p["pemantauan"]["tandaTangan"][0]["penyelia"]["nama"] == "dokumen_komitmen_tanda_tangan.jpg"


def test_eskalasi_ke_pendamping_p3h_dan_audio():
    from kawal_agent import agent as adk_agent, tools as legacy_tools
    legacy_agent = adk_agent

    c = siapkan_draf(Ctx().giliran("halo"))
    esc = T.eskalasi_ke_pendamping_p3h("Gelatin Sapi", "Belum ada sertifikat halal BPJPH", "kritis_hewani", c)
    assert esc["status"] == "berhasil_dieskalasi" and esc["ticket_id"]
    assert len(c.state.get("eskalasi_p3h") or []) == 1

    esc_legacy = legacy_tools.eskalasi_ke_pendamping_p3h("warung-001", "Gelatin Sapi", "Perlu verifikasi P3H")
    assert esc_legacy["status"] == "berhasil_dieskalasi" and esc_legacy["ticket_id"]

    lamp_legacy = legacy_tools.lampirkan_berkas("denah", "denah_dapur.jpg", c.state["pengajuan_id"])
    assert lamp_legacy["status"] == "ok" and lamp_legacy["valid"] is True

    # Batch structured output schema
    parsed = legacy_agent.AnalisisBahanResult(bahan=["Tepung Terigu", "Garam"], nomor_sertifikat=["ID1230000100101"])
    assert parsed.bahan == ["Tepung Terigu", "Garam"]

    # Multi-turn chat history fallback when DB is empty/offline
    assert isinstance(legacy_agent.build_chat_history("sesi-uji-001"), list)

    # Audio voice note native support
    meta, bagian_media = adk_agent.siapkan_lampiran(
        "+6281200000001",
        [("pesan_suara.ogg", "audio/ogg", b"OggS-dummy-voice-note")],
    )
    assert len(meta) == 1 and meta[0]["mime"] == "audio/ogg"
    assert len(bagian_media) == 1



@pytest.fixture(autouse=True)
def mock_db_offline(monkeypatch):
    monkeypatch.setitem(
        sys.modules,
        "db",
        SimpleNamespace(
            supabase=None,
            require_supabase=lambda: (_ for _ in ()).throw(RuntimeError("mock db offline")),
        ),
    )


def test_cari_rpu_halal_terdekat_dan_registrasi():
    from kawal_agent import agent as legacy_agent, tools as legacy_tools

    assert "cari_rpu_halal_terdekat" in legacy_agent.TOOL_REGISTRY
    assert "cari_rpu_halal_terdekat" in legacy_agent.SYSTEM_INSTRUCTION

    r_unggas = legacy_tools.cari_rpu_halal_terdekat("Bandung", "unggas")
    assert r_unggas["status"] == "ok"
    assert r_unggas["google_maps_url"] == "https://www.google.com/maps/search/?api=1&query=RPU+Halal+Bandung"
    assert len(r_unggas["rekomendasi"]) >= 1
    assert all(x["status_sertifikasi"] == "Berlaku" for x in r_unggas["rekomendasi"])

    r_sapi = T.cari_rpu_halal_terdekat("Pasar Ciroyom Bandung", "sapi")
    assert r_sapi["status"] == "ok"
    assert "https://www.google.com/maps/search/?api=1&query=RPU+Halal+Pasar+Ciroyom+Bandung" in r_sapi["google_maps_url"]
    assert any(x["tipe"] == "RPH" for x in r_sapi["rekomendasi"])


def test_validasi_kelayakan_sihalal_reflection_critic():
    from kawal_agent import agent as legacy_agent, tools as legacy_tools

    assert "validasi_kelayakan_sihalal" in legacy_agent.TOOL_REGISTRY

    c = siapkan_draf(Ctx().giliran("mulai"))
    # Sebelum dilengkapi: harus mendeteksi celah pada 4 alur produksi, bahan hewani, dan tim halal
    cek_awal = legacy_tools.validasi_kelayakan_sihalal(c.state["pengajuan_id"])
    assert cek_awal["layak_submit"] is False
    assert "alur_produksi" in cek_awal["bagian_kurang"]
    assert "verifikasi_bahan_hewani" in cek_awal["bagian_kurang"]
    assert "tim_halal" in cek_awal["bagian_kurang"]

    # Setelah seluruh data, 4 alur produksi, bahan hewani, tim halal, dan komitmen dilengkapi
    lengkapi_semua(c)
    setujui_pernyataan(c)
    cek_akhir = T.validasi_kelayakan_sihalal(c.state["pengajuan_id"], c)
    assert cek_akhir["layak_submit"] is True
    assert cek_akhir["audit_internal"]["alur_produksi_4_tahap"]["lengkap"] is True
    assert cek_akhir["audit_internal"]["verifikasi_bahan_hewani"]["lengkap"] is True
    assert cek_akhir["audit_internal"]["data_tim_halal"]["lengkap"] is True


def test_pindai_nota_atau_label_vision_ocr():
    from kawal_agent import agent as legacy_agent, tools as legacy_tools

    assert "pindai_nota_atau_label" in legacy_agent.TOOL_REGISTRY

    h1 = legacy_tools.pindai_nota_atau_label("nota_kecap.jpg")
    assert h1["status"] == "ok"
    assert h1["nama_bahan"] == "Kecap Manis"
    assert h1["merk"] == "Kecap Cap Layar"
    assert h1["nomor_sh_pirt"] == "ID1230000100101"
    assert h1["tanggal"]

    c = siapkan_draf(Ctx().giliran("halo"))
    h2 = T.pindai_nota_atau_label("bon_ayam_rpu.jpg", c)
    assert h2["status"] == "ok"
    assert "Ayam" in h2["nama_bahan"]
    assert h2["nomor_sh_pirt"] == "ID3175000200002"


def test_ekspor_berkas_sjph_siap_unduh():
    from pathlib import Path
    from openpyxl import load_workbook
    from kawal_agent import agent as legacy_agent

    assert "ekspor_berkas_sjph" in legacy_agent.TOOL_REGISTRY

    c = siapkan_draf(Ctx().giliran("mulai"))
    lengkapi_semua(c)
    setujui_pernyataan(c)

    h = T.ekspor_berkas_sjph(c)
    assert h["status"] == "dibuat"
    assert h["url"] and "?t=" in h["url"]
    assert h["ringkasan_sjph"]["jumlah_produk"] >= 1
    assert h["ringkasan_sjph"]["jumlah_bahan"] >= 3
    path_xlsx = Path(h["file_path"])
    assert path_xlsx.is_file()
    wb = load_workbook(path_xlsx)
    assert "Manual SJPH" in wb.sheetnames
    assert Path(h["file_path_sjph"]).is_file()
