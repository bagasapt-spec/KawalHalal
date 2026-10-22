"""Agen KawalHalal (ADK): satu LlmAgent dengan tool tulis ke mock SIHALAL, gerbang konfirmasi di kode,
dan guardrail keluaran. Dijalankan lewat jalankan() dari endpoint chat atau webhook WhatsApp.
"""
import asyncio
import concurrent.futures
import io
import logging
import os
import re
import threading
from pathlib import Path
from typing import Optional

from google.adk.agents import Agent
from google.adk.runners import Runner
from google.adk.sessions import InMemorySessionService
from google.genai import types

from agent import terapkan_guardrail

from . import gates
from .prompts import INSTRUKSI
from .tools import SEMUA_TOOL, _slug_user, folder_pengguna, url_berkas

logger = logging.getLogger("kawalhalal.agen")

APP_NAME = "kawalhalal"
MODEL = os.getenv("GEMINI_MODEL", "gemini-2.0-flash")
BATAS_TEKS_DOKUMEN = 8000
MIME_VISI = {"image/jpeg", "image/png", "image/webp", "application/pdf"}
MIME_AUDIO = {
    "audio/ogg", "audio/opus", "audio/mpeg", "audio/mp3", "audio/mp4",
    "audio/aac", "audio/amr", "audio/wav", "audio/x-wav", "audio/webm", "audio/flac",
}
DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _guardrail_keluaran(callback_context, llm_response):
    """after_model_callback: ganti teks yang menyatakan halal/pasti lolos dengan pesan aman (FR-017)."""
    content = getattr(llm_response, "content", None)
    if not content or not content.parts:
        return None
    berubah = False
    for part in content.parts:
        if getattr(part, "text", None):
            baru, terpicu = terapkan_guardrail(part.text)
            if terpicu:
                part.text = baru
                berubah = True
    if berubah:
        callback_context.state["guardrail_terpicu"] = int(callback_context.state.get("guardrail_terpicu", 0)) + 1
    return llm_response if berubah else None


def bangun_agen(model=MODEL) -> Agent:
    return Agent(
        name="kawalhalal",
        model=model,
        description="Agen pendamping pengajuan sertifikasi halal untuk pelaku usaha.",
        instruction=INSTRUKSI,
        tools=SEMUA_TOOL,
        before_agent_callback=gates.proses_giliran_baru,
        after_model_callback=_guardrail_keluaran,
        generate_content_config=types.GenerateContentConfig(temperature=0.2),
    )


root_agent = bangun_agen()
_sessions = InMemorySessionService()
_runner = Runner(agent=root_agent, app_name=APP_NAME, session_service=_sessions)
_kunci_user: dict[str, asyncio.Lock] = {}


# Semua pekerjaan agen berjalan di satu event loop khusus: kunci per pengguna dan sesi ADK terikat ke satu loop,
# sedangkan webhook memanggil dari thread BackgroundTasks dan dari loop utama uvicorn.
_loop_agen: Optional[asyncio.AbstractEventLoop] = None
_kunci_loop = threading.Lock()


def _ambil_loop_agen() -> asyncio.AbstractEventLoop:
    global _loop_agen
    with _kunci_loop:
        if _loop_agen is None or _loop_agen.is_closed():
            _loop_agen = asyncio.new_event_loop()
            threading.Thread(target=_loop_agen.run_forever, name="loop-agen-kawalhalal", daemon=True).start()
        return _loop_agen


def kirim_ke_loop_agen(coro) -> concurrent.futures.Future:
    """Jadwalkan coroutine ke loop agen tanpa menunggu. Galat dicatat ke log."""
    fut = asyncio.run_coroutine_threadsafe(coro, _ambil_loop_agen())

    def selesai(f):
        if not f.cancelled() and f.exception():
            logger.error("Tugas agen gagal: %s", f.exception(), exc_info=f.exception())

    fut.add_done_callback(selesai)
    return fut


async def di_loop_agen(coro):
    """Jalankan coroutine di loop agen dari loop mana pun (mis. endpoint FastAPI) dan tunggu hasilnya."""
    return await asyncio.wrap_future(asyncio.run_coroutine_threadsafe(coro, _ambil_loop_agen()))


async def jalankan_di_loop_agen(
    user_id: str,
    teks: str,
    lampiran: Optional[list[dict]] = None,
    audio_bytes: Optional[bytes] = None,
    audio_mime: str = "audio/ogg",
) -> dict:
    return await di_loop_agen(jalankan(user_id, teks, lampiran, audio_bytes=audio_bytes, audio_mime=audio_mime))


def ganti_model(model) -> None:
    """Untuk pengujian: ganti model agen (mis. LLM terskrip) dan mulai sesi baru."""
    global _runner, _sessions
    _sessions = InMemorySessionService()
    _runner = Runner(agent=bangun_agen(model), app_name=APP_NAME, session_service=_sessions)


def agen_aktif() -> bool:
    return bool(
        os.getenv("GEMINI_API_KEY")
        or os.getenv("GOOGLE_API_KEY")
        or os.getenv("GOOGLE_GENAI_USE_VERTEXAI")
        or os.getenv("GOOGLE_APPLICATION_CREDENTIALS")
    )


def _nama_aman(nama: str) -> str:
    dasar = os.path.basename(nama or "").replace("\\", "/").split("/")[-1]
    bersih = re.sub(r"[^A-Za-z0-9._\- ]", "_", dasar).strip()
    return bersih[:100] or "berkas"


def _teks_docx(data: bytes) -> str:
    import docx

    d = docx.Document(io.BytesIO(data))
    baris = [p.text for p in d.paragraphs if p.text.strip()]
    for t in d.tables:
        for r in t.rows:
            baris.append(" | ".join(c.text.strip() for c in r.cells if c.text.strip()))
    return "\n".join(baris)[:BATAS_TEKS_DOKUMEN]


def siapkan_lampiran(user_id: str, lampiran: list[dict]) -> tuple[list[dict], list[types.Part]]:
    """Simpan lampiran ke disk, kembalikan daftar berkas_masuk dan Part untuk model."""
    folder = folder_pengguna(user_id) / "unggahan"
    folder.mkdir(parents=True, exist_ok=True)
    daftar, parts = [], []
    for l in lampiran:
        nama, mime, data = _nama_aman(l["nama"]), (l.get("mime") or "").split(";")[0].lower(), l["data"]
        path = folder / nama
        path.write_bytes(data)
        daftar.append({"nama": nama, "mime": mime, "path": str(path), "url": url_berkas(user_id, f"unggahan/{nama}")})
        catatan = f"[Lampiran diterima: {nama} ({mime})]"
        if mime in MIME_VISI:
            parts.append(types.Part.from_text(text=catatan + " Isinya terlampir sebagai gambar/dokumen."))
            parts.append(types.Part.from_bytes(data=data, mime_type=mime))
        elif mime in MIME_AUDIO or mime.startswith("audio/"):
            parts.append(
                types.Part.from_text(
                    text=catatan + " Ini adalah pesan suara (voice note) dari pengguna. Dengarkan dan proses isinya sebagai instruksi atau jawaban pengguna."
                )
            )
            parts.append(types.Part.from_bytes(data=data, mime_type=mime))
        elif mime == DOCX_MIME or nama.lower().endswith(".docx"):
            try:
                parts.append(types.Part.from_text(text=f"{catatan} Isi dokumen:\n{_teks_docx(data)}"))
            except Exception:
                parts.append(types.Part.from_text(text=f"{catatan} Dokumen tidak terbaca."))
        elif mime == XLSX_MIME or nama.lower().endswith(".xlsx"):
            parts.append(types.Part.from_text(text=f"{catatan} Ini berkas xlsx; kalau formatnya dari KawalHalal pakai impor_berkas_xlsx."))
        else:
            parts.append(types.Part.from_text(text=f"{catatan} Format ini belum didukung."))
    return daftar, parts


async def jalankan(
    user_id: str,
    teks: str,
    lampiran: Optional[list[dict]] = None,
    audio_bytes: Optional[bytes] = None,
    audio_mime: str = "audio/ogg",
) -> dict:
    """Satu giliran percakapan. Giliran per pengguna diserialkan supaya state tidak balapan."""
    if not agen_aktif():
        return {"balasan": "Agen AI belum aktif: GEMINI_API_KEY (atau kredensial Vertex AI) belum diatur di server.", "alat": [], "aktif": False}
    sid = f"s-{_slug_user(user_id)}"
    async with _kunci_user.setdefault(user_id, asyncio.Lock()):
        sesi = await _sessions.get_session(app_name=APP_NAME, user_id=user_id, session_id=sid)
        if sesi is None:
            sesi = await _sessions.create_session(app_name=APP_NAME, user_id=user_id, session_id=sid, state={})
        parts = []
        delta = {}
        daftar_lampiran = list(lampiran or [])
        if audio_bytes:
            ext = (audio_mime or "audio/ogg").split("/")[-1].split(";")[0] or "ogg"
            daftar_lampiran.append({"nama": f"voicenote.{ext}", "mime": audio_mime or "audio/ogg", "data": audio_bytes})
        if daftar_lampiran:
            baru, parts_lampiran = siapkan_lampiran(user_id, daftar_lampiran)
            delta["berkas_masuk"] = (sesi.state.get("berkas_masuk") or []) + baru
            parts.extend(parts_lampiran)
        parts.insert(0, types.Part.from_text(text=teks or "(pengguna hanya mengirim lampiran)"))
        pesan = types.Content(role="user", parts=parts)

        balasan, alat = [], []
        try:
            async for e in _runner.run_async(user_id=user_id, session_id=sid, new_message=pesan, state_delta=delta or None):
                if e.content and e.content.parts:
                    for p in e.content.parts:
                        if getattr(p, "function_call", None):
                            alat.append(p.function_call.name)
                if e.is_final_response() and e.content and e.content.parts:
                    balasan.append("".join(p.text or "" for p in e.content.parts if getattr(p, "text", None)))
        except Exception as err:
            logger.error("Agen gagal: %s", err, exc_info=True)
            return {"balasan": "Maaf, agen sedang bermasalah. Coba kirim ulang pesan Anda sebentar lagi.", "alat": alat, "aktif": True, "galat": f"{type(err).__name__}: {err}"}
    return {"balasan": "\n".join(b for b in balasan if b).strip() or "Maaf, saya belum bisa menjawab. Coba ulangi pesan Anda.", "alat": alat, "aktif": True}


async def lihat_state(user_id: str) -> dict:
    sesi = await _sessions.get_session(app_name=APP_NAME, user_id=user_id, session_id=f"s-{_slug_user(user_id)}")
    return dict(sesi.state) if sesi else {}


async def hapus_sesi(user_id: str) -> bool:
    sid = f"s-{_slug_user(user_id)}"
    ada = await _sessions.get_session(app_name=APP_NAME, user_id=user_id, session_id=sid) is not None
    if ada:
        await _sessions.delete_session(app_name=APP_NAME, user_id=user_id, session_id=sid)
    return ada
